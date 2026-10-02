"""The Excel export of the code sheet: every determination lands in the workbook, numbers stay numbers,
and the server writes it on save and serves it for the massing on screen."""
import pathlib, sys
import pytest
ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / "data" / "projects")]
pytest.importorskip("openpyxl")
from openpyxl import load_workbook
from courtyard_commons import project
from codesheet import sheet
from codesheet.xlsx_sheet import to_xlsx, FLAT_COLS, DRAFT_TEXT


@pytest.fixture(scope="module")
def sd():
    return sheet.run_all(project, chosen={"North": "3.2.2.51", "South": "3.2.2.51"})


def test_workbook_has_the_three_sheets_and_every_line(tmp_path, sd):
    p = to_xlsx(sd, tmp_path / "cc.xlsx", sheet_no="MS-01", revision="test")
    wb = load_workbook(p)
    assert wb.sheetnames == ["Code summary", "Determinations", "Flags & notes"]
    flat = wb["Determinations"]
    assert [c.value for c in flat[1]] == [n for n, _ in FLAT_COLS]
    n_lines = sum(len(s.dets) for s in sd.sections)
    assert flat.max_row == n_lines + 1
    assert flat.freeze_panes == "A2" and flat.auto_filter.ref.startswith("A1:")
    # the summary sheet carries the title block, the draft stamp and each section banner
    summ = wb["Code summary"]
    col_a = [summ.cell(row=r, column=1).value for r in range(1, summ.max_row + 1)]
    assert col_a[0] == f"{project.project_name} — Building Code Analysis"
    assert DRAFT_TEXT in col_a
    for s in sd.sections:
        assert s.number in col_a, s.title
    assert "MS-01 · test" in [summ.cell(row=r, column=2).value for r in range(1, 14)]


def test_values_stay_numeric_and_clauses_cite_pages(tmp_path, sd):
    wb = load_workbook(to_xlsx(sd, tmp_path / "cc.xlsx"))
    flat = wb["Determinations"]
    rows = {r[4]: r for r in flat.iter_rows(min_row=2, values_only=True)}        # keyed by the determination key
    area = next(k for k in rows if k.endswith("building_area_m2") and k.startswith("North"))
    assert isinstance(rows[area][5], (int, float)) and rows[area][6] == "m²"
    article = next(k for k in rows if k == "North.article")
    assert rows[article][5] == "3.2.2.51"
    assert rows[article][8] and rows[article][10]                               # clause id and PDF page
    # the per-storey occupant-load appendix rows are numbers of persons
    appx = [r for r in rows.values() if r[0] == "C"]
    assert appx and all(isinstance(r[5], (int, float)) for r in appx)


def test_flags_and_not_covered_are_listed(tmp_path, sd):
    wb = load_workbook(to_xlsx(sd, tmp_path / "cc.xlsx"))
    ws = wb["Flags & notes"]
    cells = [ws.cell(row=r, column=c).value for r in range(1, ws.max_row + 1) for c in (1, 2)]
    assert f"Reviewer flags ({len(sd.flags)})" in cells
    for _, text in sd.flags:
        assert text in cells
    for text in sheet.NOT_COVERED:
        assert text in cells


def test_server_writes_the_workbook_on_save_and_builds_one_on_demand(tmp_path, monkeypatch):
    pytest.importorskip("fastapi")
    from app import server
    from codesheet.massing import MassingSpec
    monkeypatch.setattr(server, "PROJECTS", tmp_path / "projects")
    spec = MassingSpec.model_validate(server.example_demo())
    rec = server.save_iteration(server.SaveRequest(spec=spec, name="xlsx test", note=""))
    assert rec["has_xlsx"] is True
    d = server._find(rec["id"])
    assert (d / "summary.xlsx").exists() and (d / "summary.html").exists() and not (d / "summary.pdf").exists()
    assert load_workbook(d / "summary.xlsx")["Code summary"].cell(row=11, column=2).value == "MS-01 · Massing iteration — xlsx test"
    resp = server.summary_xlsx(spec, name="on screen")
    assert resp.media_type == server.XLSX_MEDIA and resp.body[:2] == b"PK"
    assert 'filename="courtyard-commons-example-code-summary.xlsx"' in resp.headers["content-disposition"]
